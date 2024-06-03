import io
import pandas as pd

import google.auth
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload
from dotenv import load_dotenv

from docx import Document
import openpyxl
import csv


load_dotenv()

class GoogleDriveImporter:
  def __init__(self):
    load_dotenv()
    self.creds, _ = google.auth.default()
    self.service = build("drive", "v3", credentials=self.creds)

  def download_file(self, real_file_id) -> io.BytesIO:
    """Downloads a file
    Args:
      real_file_id: ID of the file to download
    Returns : IO object with location.
    """
    try:
      file_id = real_file_id
      request = self.service.files().get_media(fileId=file_id)
      file = io.BytesIO()
      downloader = MediaIoBaseDownload(file, request)
      done = False
      while done is False:
        status, done = downloader.next_chunk()
        # print(f"Download {int(status.progress() * 100)}.")
    except HttpError as error:
      print(f"An error occurred: {error}")
      file = None
    return file

  def read_docx(self, file_content):
    doc = Document(io.BytesIO(file_content))
    for para in doc.paragraphs:
      print(para.text)

  def read_xlsx(self, file_content):
    workbook = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    worksheet = workbook.active

    for row in worksheet.iter_rows(values_only=True):
      row_values = '\t'.join([str(cell) if cell is not None else '' for cell in row])
      print(row_values)

  def read_csv(self, file_content) -> pd.DataFrame:
    reader = csv.reader(io.StringIO(file_content.decode("utf-8")), delimiter=';')
    df = pd.DataFrame(reader)
    df.columns = df.iloc[0]
    df = df[1:]
    return df


if __name__ == "__main__":
  importer = GoogleDriveImporter()

  # importer.read_docx(docx_output)

  # importer.read_xlsx(sheets_output)
  csv_output = importer.download_file(real_file_id="asdfasdf")
  print(importer.read_csv(csv_output).head())
